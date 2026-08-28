# -*- coding: utf-8 -*-
"""
Excel 导出模块：明细按「采集时段 × 有/无早餐」分 Sheet；另有城市汇总 / 趋势图 / 波动率。
"""
import io
import re
from datetime import datetime
from pathlib import Path

import analyzer
import matplotlib

matplotlib.use("Agg")  # 无界面后端，避免 GUI 弹窗
import matplotlib.pyplot as plt
from matplotlib import font_manager

import config
import storage

# ------------------------------------------------------------------
# 中文字体配置（Windows 常见中文字体）
# ------------------------------------------------------------------
def _setup_cn_font():
    candidates = ["Microsoft YaHei", "SimHei", "SimSun", "PingFang SC", "Noto Sans CJK SC"]
    for name in candidates:
        try:
            font_manager.findfont(name, fallback_to_default=False)
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["axes.unicode_minus"] = False
            return
        except Exception:
            continue
    # 兜底：尝试注册本机字体
    for fp in [
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]:
        p = Path(fp)
        if p.exists():
            font_manager.fontManager.addfont(str(p))
            name = font_manager.FontProperties(fname=str(p)).get_name()
            plt.rcParams["font.sans-serif"] = [name]
            plt.rcParams["axes.unicode_minus"] = False
            return
    print("警告：未找到中文字体，趋势图中文可能显示为方框")


_setup_cn_font()

# ------------------------------------------------------------------
# 导出路径（统一归纳到 输出/excel/YYYY-MM-DD/）
# ------------------------------------------------------------------
_WIN_BAD = re.compile(r'[\\/:*?"<>|\x00-\x1f\x7f]')


def _safe_filename(text, max_len=30):
    """去掉 Windows 文件名非法字符（含换行等控制符）并截断。"""
    if not text:
        return ""
    text = _WIN_BAD.sub("_", str(text))
    text = re.sub(r"\s+", " ", text).strip().strip(".")
    if max_len:
        text = text[:max_len].rstrip(" .")
    return text or "export"


def _keyword_slug(keyword, max_len=24):
    """多酒店关键词压成合法短文件名：单家用店名，多家用「第一家等N家」。"""
    if not keyword:
        return ""
    names = [n.strip() for n in re.split(r"[\n\r,，;；|]+", str(keyword)) if n.strip()]
    if not names:
        return ""
    if len(names) == 1:
        return _safe_filename(names[0], max_len)
    suffix = f"等{len(names)}家"
    head = _safe_filename(names[0], max(4, max_len - len(suffix)))
    return f"{head}{suffix}"


def resolve_export_path(output_path=None, *, city=None, start_date=None, end_date=None, keyword=None):
    """
    解析最终导出路径。
    - 未指定路径：输出/excel/今天/城市_日期范围_时间.xlsx
    - 仅文件名：输出/excel/今天/文件名.xlsx
    - 绝对路径：按用户指定
    - 含子目录的相对路径：相对于 输出/ 目录
    """
    today_dir = config.OUTPUT_EXCEL_DIR / datetime.now().strftime("%Y-%m-%d")
    today_dir.mkdir(parents=True, exist_ok=True)

    if output_path:
        p = Path(output_path)
        if p.suffix.lower() != ".xlsx":
            p = p.with_suffix(".xlsx")
        if p.is_absolute():
            p.parent.mkdir(parents=True, exist_ok=True)
            return p
        if len(p.parts) == 1:
            stem = _safe_filename(p.stem, max_len=0)
            return today_dir / f"{stem}.xlsx"
        out = config.OUTPUT_DIR / p
        out.parent.mkdir(parents=True, exist_ok=True)
        return out

    parts = []
    if city:
        parts.append(_safe_filename(city, 20))
    slug = _keyword_slug(keyword, 24)
    if slug:
        parts.append(slug)
    if start_date:
        s = start_date.replace("-", "")
        e = (end_date or start_date).replace("-", "")
        parts.append(s if s == e else f"{s}-{e}")
    parts.append(datetime.now().strftime("%H%M%S"))
    if parts:
        name = "_".join(parts) + ".xlsx"
    else:
        name = config.EXPORT_FILENAME_TEMPLATE.format(
            date=datetime.now().strftime("%Y%m%d_%H%M%S")
        )
    return today_dir / name


# ------------------------------------------------------------------
# 明细按采集时段 × 早餐分 Sheet
# ------------------------------------------------------------------
def _row_time_slot(row) -> str:
    if "time_slot" in row.keys() and row["time_slot"]:
        return str(row["time_slot"])
    crawled = row["crawled_at"] if "crawled_at" in row.keys() else None
    if crawled:
        try:
            dt = datetime.strptime(str(crawled)[:19], "%Y-%m-%d %H:%M:%S")
            return storage.infer_time_slot(dt)
        except ValueError:
            pass
    return storage.TIME_SLOTS[0]


def _row_breakfast(row) -> str:
    room = row["room_type"] if "room_type" in row.keys() else ""
    return storage.classify_breakfast(room)


def _slot_hour_label(slot: str) -> str:
    hour_part = str(slot).split(":")[0]
    try:
        return str(int(hour_part))
    except ValueError:
        return hour_part or "0"


def _detail_sheet_title(slot: str, breakfast: str) -> str:
    """Excel 工作表名不能含冒号，如「明细_10点_有早餐」。"""
    if slot == "其他":
        return f"明细_其他_{breakfast}"
    return f"明细_{_slot_hour_label(slot)}点_{breakfast}"


def _group_rows_by_slot_breakfast(rows) -> list[tuple[str, str, list]]:
    """
    按 时段×早餐 分桶，返回有序列表 [(slot, breakfast, rows), ...]。
    标准四档 × 有/无早餐始终出现（可为空表）；未知时段归「其他」。
    """
    buckets: dict[tuple[str, str], list] = {}
    for slot in storage.TIME_SLOTS:
        for bf in storage.BREAKFAST_LABELS:
            buckets[(slot, bf)] = []
    other: dict[str, list] = {bf: [] for bf in storage.BREAKFAST_LABELS}

    for r in rows:
        slot = _row_time_slot(r)
        bf = _row_breakfast(r)
        if slot in storage.TIME_SLOTS:
            buckets[(slot, bf)].append(r)
        else:
            other[bf].append(r)

    ordered = [(s, b, buckets[(s, b)]) for s in storage.TIME_SLOTS for b in storage.BREAKFAST_LABELS]
    for bf in storage.BREAKFAST_LABELS:
        if other[bf]:
            ordered.append(("其他", bf, other[bf]))
    return ordered


def _write_detail_sheet(ws, rows):
    headers = ["城市", "酒店名称", "入住日期", "房型", "早餐", "价格(元/晚)", "采集时段", "来源", "采集时间"]
    ws.append(headers)
    for r in rows:
        ws.append([
            r["city"],
            r["hotel_name"],
            r["checkin_date"],
            r["room_type"],
            _row_breakfast(r),
            r["price"],
            _row_time_slot(r),
            r["source"],
            r["crawled_at"],
        ])
    _style_header(ws, headers)
    if ws.max_row >= 1 and ws.max_column >= 1:
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"


# ------------------------------------------------------------------
# 图表颜色（浅色背景 + 可读深色线条，每城市一色）
# ------------------------------------------------------------------
TREND_COLORS = [
    "#185FA5", "#0F6E56", "#993C1D", "#854F0B",
    "#534AB7", "#993556", "#3B6D11", "#A32D2D",
]


def make_trend_chart(trend):
    """生成价格趋势折线图，返回 PNG 字节。trend = {city: [(date, mean, n)]}"""
    fig, ax = plt.subplots(figsize=(10, 5.5), dpi=110)
    for i, (city, series) in enumerate(trend.items()):
        dates = [d for d, _, _ in series]
        means = [m for _, m, _ in series]
        ax.plot(dates, means, marker="o", linewidth=1.8, markersize=5,
                color=TREND_COLORS[i % len(TREND_COLORS)], label=city)

    ax.set_title("各城市酒店日均价趋势", fontsize=14, fontweight="bold", pad=14)
    ax.set_xlabel("入住日期")
    ax.set_ylabel("平均房价（元/晚）")
    ax.grid(True, linestyle="--", alpha=0.35)
    ax.legend(loc="best", frameon=False)

    fig.autofmt_xdate()
    fig.tight_layout()

    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf.getvalue()


def export_excel(rows, output_path=None, trend=None, volatility=None, mom=None,
                 *, city=None, start_date=None, end_date=None, keyword=None):
    """导出多 Sheet Excel。明细按「时段×早餐」分 Sheet；汇总/趋势/波动基于全部 rows。"""
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font

    output_path = resolve_export_path(
        output_path, city=city, start_date=start_date, end_date=end_date, keyword=keyword,
    )

    trend = trend if trend is not None else analyzer.price_trend(rows)
    volatility = volatility if volatility is not None else analyzer.hotel_volatility(rows)
    mom = mom if mom is not None else analyzer.price_mom(rows)

    wb = Workbook()
    groups = (
        _group_rows_by_slot_breakfast(rows)
        if rows
        else [(s, b, []) for s in storage.TIME_SLOTS for b in storage.BREAKFAST_LABELS]
    )

    # ---------- 明细：时段 × 有/无早餐 ----------
    first = True
    for slot, breakfast, slot_rows in groups:
        title = _detail_sheet_title(slot, breakfast)
        if first:
            ws = wb.active
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title)
        _write_detail_sheet(ws, slot_rows)

    # ---------- 城市汇总 ----------
    ws2 = wb.create_sheet("城市汇总")
    h2 = ["城市", "入住日期", "样本数", "均价(元)", "中位数", "最低价", "最高价", "标准差", "变异系数CV"]
    ws2.append(h2)
    summary = analyzer.city_date_summary(rows)
    for s in summary:
        ws2.append([s["city"], s["date"], s["count"], s["mean"], s["median"],
                    s["min"], s["max"], s["std"], s["cv"]])
    _style_header(ws2, h2)

    # ---------- 价格趋势图 ----------
    ws3 = wb.create_sheet("价格趋势")
    if trend:
        png = make_trend_chart(trend)
        img = XLImage(io.BytesIO(png))
        img.width = 900
        img.height = 500
        ws3.add_image(img, "A1")
        ws3["A35"] = "说明：折线为各城市酒店日均价（每个入住日期的均价）"
        ws3["A35"].font = Font(size=10, color="666666")

    # ---------- 波动率排行 ----------
    ws4 = wb.create_sheet("波动率排行")
    h4 = ["排名", "城市", "酒店名称", "有效日期数", "均价(元)", "标准差", "变异系数CV", "最低价", "最高价"]
    ws4.append(h4)
    for i, v in enumerate(volatility, 1):
        ws4.append([i, v["city"], v["hotel_name"], v["count"], v["mean"],
                    v["std"], v["cv"], v["min"], v["max"]])
    _style_header(ws4, h4)
    ws4.append([])
    ws4.append(["环比说明", "相邻采集日期的均价变化率"])
    ws4["A12"] = "环比变化"
    ws4["A12"].font = Font(bold=True, size=12)
    for i, m in enumerate(mom[:30], 1):
        ws4.append([f"  {m['city']}", f"{m['from_date']} → {m['to_date']}",
                    f"均价 {m['prev_mean']} → {m['cur_mean']}",
                    f"{m['change_pct']}%"])

    wb.save(output_path)
    return output_path


def _style_header(ws, headers):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    fill = PatternFill("solid", fgColor="185FA5")
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    # 自动列宽（粗略）
    for col_idx in range(1, len(headers) + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            v = row[0].value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        ws.column_dimensions[letter].width = min(max_len + 4, 60)
